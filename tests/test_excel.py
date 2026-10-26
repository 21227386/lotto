from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from excel_io import HEADERS, LEGACY_HEADERS, read_draws, read_workbook, write_analysis, write_template
from model import estimate, recommend


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / 'input.xlsx'

    def tearDown(self):
        self.tmp.cleanup()

    def workbook(self, bad=None, header=True, with_ids=True):
        book = Workbook()
        sheet = book.active
        if header:
            sheet.append(LEGACY_HEADERS if with_ids else LEGACY_HEADERS[1:])
        for i in range(412):
            row = ([836 + i] if with_ids else []) + [1, 8, 15, 22, 29, 36]
            if i == 0 and bad is not None:
                row = bad
            sheet.append(row)
        book.save(self.path)
        book.close()

    def test_read_and_export(self):
        self.workbook()
        draws, ids, _ = read_draws(self.path)
        self.assertEqual((len(draws), min(ids), max(ids)), (412, 836, 1247))
        estimates = estimate(draws)
        numbers = recommend(estimates)
        output = Path(self.tmp.name) / 'result.xlsx'
        write_analysis(output, estimates, numbers, ids, 412, False)
        book = load_workbook(output)
        self.assertEqual(book['추천결과']['B2'].value, 1248)
        self.assertEqual(list(next(book['추천결과'].iter_rows(min_row=6, max_row=6, min_col=2, max_col=7, values_only=True))), numbers)
        self.assertEqual(book['번호별분석'].max_row, 46)
        book.close()

    def test_six_columns_and_no_header(self):
        self.workbook(header=False, with_ids=False)
        draws, ids, _ = read_draws(self.path)
        self.assertEqual(len(draws), 412)
        self.assertEqual(ids, [])

    def test_bad_cells_and_extra_columns(self):
        for bad in [[836, 1, 8, 15, 22, 29, '=30+6'],
                    [836, 1, 8, 15, 22, 29, None],
                    [836, True, 8, 15, 22, 29, 36],
                    [836, 1.5, 8, 15, 22, 29, 36],
                    [836, 1, 8, 15, 22, 29, 36, 7]]:
            with self.subTest(bad=bad):
                self.workbook(bad)
                with self.assertRaises(ValueError):
                    read_draws(self.path)

    def test_template_and_invalid_files(self):
        write_template(self.path)
        book = load_workbook(self.path)
        self.assertEqual(list(next(book.active.values)), HEADERS)
        book.close()
        with self.assertRaises(ValueError):
            read_draws(self.path)
        self.path.write_bytes(b'not excel')
        with self.assertRaises(ValueError):
            read_draws(self.path)
        with self.assertRaises(ValueError):
            read_draws(self.path.with_suffix('.xls'))

    def modern_workbook(self, count=413, reverse=False, bad_date=False, missing=False):
        from datetime import datetime, timedelta
        book = Workbook()
        sheet = book.active
        sheet.title = 'Lotto Draws'
        sheet.append(HEADERS)
        rows = [[836 + i, datetime(2018, 12, 8) + timedelta(weeks=i), 1, 8, 15, 22, 29, 36] for i in range(count)]
        if bad_date:
            rows[-1][1] = '=TODAY()'
        if missing:
            rows[-2][0] = 9999
        for row in reversed(rows) if reverse else rows:
            sheet.append(row)
        book.save(self.path)
        book.close()

    def test_all_rows_and_dates_are_retained(self):
        self.modern_workbook(reverse=True)
        data = read_workbook(self.path)
        self.assertEqual(data.source_count, 413)
        self.assertEqual((data.ids[0], data.ids[-1]), (836, 1248))
        self.assertEqual(data.dates[-1].isoformat(), '2026-10-31')
        self.assertEqual(len(data.draws), 413)

    def test_modern_rejects_bad_dates_and_missing_draws(self):
        for options in [{'bad_date': True}, {'missing': True}, {'count': 411}]:
            with self.subTest(options=options):
                self.modern_workbook(**options)
                with self.assertRaises(ValueError):
                    read_workbook(self.path)

    def test_date_text_and_date_order(self):
        self.modern_workbook(count=412)
        book = load_workbook(self.path)
        book.active['B413'] = '2026-10-24'
        book.save(self.path)
        book.close()
        self.assertEqual(read_workbook(self.path).dates[-1].isoformat(), '2026-10-24')
        book = load_workbook(self.path)
        book.active['B413'] = '2018-12-08'
        book.save(self.path)
        book.close()
        with self.assertRaises(ValueError):
            read_workbook(self.path)


if __name__ == '__main__':
    unittest.main()
