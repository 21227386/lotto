"""Real Tk integration check; requires a display. Uses synthetic data only."""
from pathlib import Path
import random
from datetime import datetime, timedelta
import tempfile
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from excel_io import HEADERS
from lotto_gui import App
from model import COMPOSITION_PLAN, SUBCOMPOSITION_PLAN, estimate, frequency_groups, frequency_subgroups


def classes(widget):
    return [widget.winfo_class()] + [name for child in widget.winfo_children() for name in classes(child)]


def run():
    with tempfile.TemporaryDirectory() as tmp:
        source = Path(tmp) / 'input.xlsx'
        output = Path(tmp) / 'output.xlsx'
        book = Workbook()
        book.active.append(HEADERS)
        rng = random.Random(43)
        for draw in range(836, 1249):
            book.active.append([draw, datetime(2018, 12, 8) + timedelta(weeks=draw - 836)] + rng.sample(range(1, 46), 6))
        book.save(source)
        book.close()
        app = App()
        errors = []
        try:
            app.update()
            with patch('lotto_gui.messagebox.showerror', side_effect=lambda *args: errors.append(args)):
                app.load_path(source)
                expected_groups = frequency_subgroups(estimate(app.data.draws))
                for name, variable, parent in zip(('고빈도', '중빈도', '저빈도'), app.group_vars, expected_groups):
                    assert variable.get() == name + ': ' + '  |  '.join(' '.join(f'{n:02d}' for n in part) for part in parent)
                app.analyze()
                app.update()
                assert app.title() == 'Lotto Generator'
                assert app.copyright_label['text'] == 'Hun-Kee, Lee All rights Reserved.'
                assert 'Times New Roman' in str(app.copyright_label['font'])
                assert app.copyright_label.winfo_ismapped()
                assert app.copyright_label.winfo_y() + app.copyright_label.winfo_height() <= app.winfo_height()
                assert app.select_button['text'] == 'Select Input File'
                assert 'Text' not in classes(app) and 'Treeview' not in classes(app)
                assert len(app.recommendations) == 5
                assert len({r.numbers for r in app.recommendations}) == 5
                assert all(len(set(r.numbers)) == 6 for r in app.recommendations)
                groups = frequency_groups(estimate(app.data.draws))
                assert tuple(r.composition for r in app.recommendations) == COMPOSITION_PLAN
                assert tuple(r.subcomposition for r in app.recommendations) == SUBCOMPOSITION_PLAN
                assert all(tuple(tuple(len(set(r.numbers) & set(part)) for part in parent) for parent in expected_groups)
                           == r.subcomposition for r in app.recommendations)
                assert all(tuple(len(set(r.numbers) & set(group)) for group in groups) == r.composition
                           for r in app.recommendations)
                assert '1249' in app.status.get() and '413' in app.status.get()
                assert app.data.ids[0] == 836 and len(app.data.draws) == 413
                assert app.header_image.width() > 0 and app.banner.winfo_ismapped()
                first = list(app.recommendations)
                with patch('lotto_gui.balanced_recommendations', side_effect=RuntimeError('cached result should be reused')):
                    app.analyze()
                assert app.recommendations == first
                assert [[label['text'] for label in row] for row in app.number_labels] == [
                    [f'{n:02d}' for n in result.numbers] for result in first]
                with patch('lotto_gui.filedialog.asksaveasfilename', return_value=str(output)):
                    app.save()
                result = load_workbook(output)
                assert result['설정']['B2'].value == 1249
                assert result['설정']['B3'].value == 413
                assert result['추천번호'].max_row == 6
                assert list(result['추천번호'].iter_rows(min_row=2, values_only=True)) == [
                    (i + 1,) + r.numbers + (':'.join(map(str, r.composition)), ', '.join(':'.join(map(str, p)) for p in r.subcomposition)) for i, r in enumerate(first)]
                assert result['빈도그룹'].max_row == 46
                result.close()
                app.load_path(source)
                assert not app.recommendations
                with patch('lotto_gui.balanced_recommendations', return_value=first) as fresh:
                    app.analyze()
                    fresh.assert_called_once()
                app.prior_strength = float('nan')
                app.analyze()
                assert not app.recommendations and len(errors) == 1
                app.prior_strength = 412
                app.analyze()
                try:
                    app.load_path(Path(tmp) / 'missing.xlsx')
                except ValueError:
                    pass
                assert app.data is None and not app.recommendations
                assert all(v.get().endswith('—') for v in app.group_vars)
                assert all(label['text'] == '—' for row in app.number_labels for label in row)
            print('GUI_SMOKE_OK: all 413 rows, compact GUI, nested nine-group quotas, ranked group labels, font and copyright, Excel export, invalidation')
        finally:
            app.destroy()


if __name__ == '__main__':
    run()
