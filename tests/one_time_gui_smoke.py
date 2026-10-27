import hashlib
from pathlib import Path
import tempfile
from unittest.mock import patch
from openpyxl import load_workbook
from one_time_gui import OneTimeApp, PasswordWindow, SUBTITLE
from one_time_data import load_embedded_data
from one_time_access import LaunchLimiter
from model import passes_pattern_rules, SUBCOMPOSITION_PLAN


def texts(widget):
    found = []
    if 'text' in widget.keys():
        found.append(str(widget.cget('text')))
    for child in widget.winfo_children():
        found.extend(texts(child))
    return found


with tempfile.TemporaryDirectory() as temp, patch('one_time_access.PASSWORD_SHA256', hashlib.sha256(b'12345').hexdigest()):
    limiter = LaunchLimiter(Path(temp) / 'usage.sqlite3')
    login = PasswordWindow(limiter)
    with patch('one_time_gui.messagebox.showerror') as error:
        login.password.set('00000')
        login.submit()
        assert error.call_count == 1 and limiter.used() == 0
    login.password.set('12345')
    login.submit()
    assert login.granted == 1 and limiter.used() == 1
    app = OneTimeApp(load_embedded_data(), 1)
    try:
        app.update()
        assert app.subtitle_label['text'] == SUBTITLE
        assert len(app.number_labels) == 5 and app.machine.set_count == 5
        assert 'Select Input File' not in texts(app)
        assert '입력용 엑셀 양식 저장' not in texts(app)
        try:
            app.load_path('anything.xlsx')
            raise AssertionError('external input must be disabled')
        except PermissionError:
            pass
        app.analyze()
        assert len(app.recommendations) == 5
        assert tuple(r.subcomposition for r in app.recommendations) == SUBCOMPOSITION_PLAN
        assert all(passes_pattern_rules(r.numbers) for r in app.recommendations)
        app.finish_reveal()
        assert app.revealed_count == 30
        assert 'ALL 5 SETS COMPLETE' == app.machine.canvas.itemcget(app.machine.caption, 'text')
        path = Path(temp) / 'results.xlsx'
        with patch('one_time_gui.filedialog.asksaveasfilename', return_value=str(path)):
            app.save()
        book = load_workbook(path)
        assert book['추천번호'].max_row == 6
        assert book['설정']['B2'].value == 1248
        assert ('추천 세트 수', 5) in list(book['설정'].values)
        book.close()
    finally:
        app.destroy()
print('ONE_TIME_GUI_OK: wrong password uncharged, correct password, embedded data, no input UI, five valid sets and fixed export')
