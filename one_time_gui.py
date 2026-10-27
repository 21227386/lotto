"""Separate fixed-data, password-protected, two-launch Windows application."""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter import font as tkfont

from lotto_gui import App
from model import COMPOSITION_PLAN, SUBCOMPOSITION_PLAN, estimate, balanced_recommendations
from excel_io import write_recommendations
from one_time_data import load_embedded_data
from one_time_access import LaunchLimiter, LaunchLimitReached, AccessStateError, BLOCKED_MESSAGE, LicenseExpired, ensure_not_expired, EXPIRED_MESSAGE
from openpyxl import load_workbook

APP_TITLE = 'One-time Lotto Generator for You'
SUBTITLE = '1248회 당첨번호 예측'


def write_fixed_results(path, results, data, strength):
    write_recommendations(path, results, data, strength)
    book = load_workbook(path)
    sheet = book['설정']
    sheet.delete_rows(1, sheet.max_row)
    for row in [('항목', '값'), ('대상 회차', 1248), ('분석 회차 수', 412), ('분석 범위', '836~1247'),
                ('내장 데이터', '외부 입력 없음'), ('사전 강도', strength), ('추천 세트 수', 5),
                ('고:중:저 구성', '2:2:2 2세트 / 3:1:2 2세트 / 2:1:3 1세트'),
                ('번호 제한', '연속 3개 금지 / n,n+7,n+14 금지'),
                ('주의', '실제 당첨 우위는 입증되지 않았습니다.')]:
        sheet.append(row)
    for index, plan in enumerate(SUBCOMPOSITION_PLAN, 1):
        sheet.append([f'SET {index} 세부 구성', ', '.join(':'.join(map(str, part)) for part in plan)])
    book.save(path)
    book.close()


class OneTimeApp(App):
    def __init__(self, data, launch_number):
        ensure_not_expired()
        self._expiry_timer = None
        super().__init__()
        self.title(APP_TITLE)
        self.geometry('1220x850')
        self.minsize(1180, 820)
        self.machine.set_count = 5
        self.launch_number = launch_number
        outer = self.banner.master
        title = next(w for w in outer.winfo_children() if isinstance(w, ttk.Label) and w.cget('text') == 'Lotto Generator')
        self.subtitle_label = ttk.Label(outer, text=SUBTITLE, font=('Times New Roman', 15, 'bold'))
        self.subtitle_label.pack(after=title, anchor='w', pady=(0, 8))
        controls = self.select_button.master
        for child in list(controls.winfo_children()):
            if child is not self.generate_button:
                child.destroy()
        for child in list(self.save_button.master.winfo_children()):
            if isinstance(child, ttk.Button) and child.cget('text') == '입력용 엑셀 양식 저장':
                child.destroy()
        for row in self.number_labels[5:]:
            row[0].master.destroy()
        self.number_labels = self.number_labels[:5]
        for index, row in enumerate(self.number_labels):
            label = next(w for w in row[0].master.winfo_children() if isinstance(w, ttk.Label))
            label.configure(text=f'SET {index + 1}  ({":".join(map(str, COMPOSITION_PLAN[index]))})')
        for child in outer.winfo_children():
            if isinstance(child, ttk.Label) and child.cget('text').startswith('동률은 무작위 선택'):
                child.configure(text=f'내장 412회 전용 · 1248회 예측 · 이번 실행 {launch_number}/2 · 사용기한 2026-10-31')
        self.generate_button.configure(text='분석 후 5세트 추출')
        self.save_button.configure(text='추천 5세트 엑셀 저장')
        self.data = data
        self.show_groups()
        self.status.set(f'836~1247회 내장 412회 · 1248회 예측 준비 · 실행 {launch_number}/2')
        self._expiry_timer = self.after(10_000, self._watch_expiry)

    def _watch_expiry(self):
        self._expiry_timer = None
        if self._allow_action():
            self._expiry_timer = self.after(10_000, self._watch_expiry)

    def _allow_action(self):
        try:
            ensure_not_expired()
            return True
        except LicenseExpired:
            self.cancel_reveal()
            messagebox.showinfo('사용기한 만료', EXPIRED_MESSAGE, parent=self)
            self.destroy()
            return False

    def destroy(self):
        if self._expiry_timer is not None:
            self.after_cancel(self._expiry_timer)
            self._expiry_timer = None
        super().destroy()

    def _on_destroy(self, event):
        if event.widget is self and self._expiry_timer is not None:
            self.after_cancel(self._expiry_timer)
            self._expiry_timer = None
        super()._on_destroy(event)

    def load(self):
        raise PermissionError('이 프로그램은 내장 데이터만 사용합니다.')

    def load_path(self, path):
        raise PermissionError('외부 입력 파일을 사용할 수 없습니다.')

    def template(self):
        raise PermissionError('입력 양식을 제공하지 않는 내장 데이터 전용 프로그램입니다.')

    def analyze(self):
        if not self._allow_action():
            return
        if self.revealing or (self.recommendations and self.generated_strength == self.prior_strength):
            return
        self.clear_results()
        try:
            # One highest valid result from each of the original five profiles.
            self.recommendations = balanced_recommendations(estimate(self.data.draws, self.prior_strength))[::2]
            self.generated_strength = self.prior_strength
        except ValueError as exc:
            messagebox.showerror('분석 실패', str(exc), parent=self)
            return
        self.revealing = True
        self.generate_button.configure(state='disabled')
        self.skip_button.configure(state='normal')
        self.machine.start()
        self.status.set('1248회 예측 · 추첨 연출 중 · SET 1 준비')
        self._reveal_timer = self.after(self.WARMUP_MS, self.reveal_next)

    def reveal_next(self):
        if self._allow_action():
            super().reveal_next()

    def finish_reveal(self):
        if self._allow_action():
            super().finish_reveal()

    def save(self):
        if not self._allow_action():
            return
        if self.revealing or not self.recommendations:
            messagebox.showinfo('결과 표시 필요', '5세트 추출을 완료하거나 즉시 표시를 눌러 주세요.', parent=self)
            return
        path = filedialog.asksaveasfilename(defaultextension='.xlsx', initialfile='Lotto_1248_5sets.xlsx',
                                          filetypes=[('Excel 통합 문서', '*.xlsx')])
        if path and self._allow_action():
            try:
                write_fixed_results(path, self.recommendations, self.data, self.prior_strength)
            except OSError as exc:
                messagebox.showerror('저장 실패', str(exc), parent=self)


class PasswordWindow(tk.Tk):
    def __init__(self, limiter):
        super().__init__()
        self.title(APP_TITLE + ' - 비밀번호')
        self.geometry('430x230')
        self.resizable(False, False)
        self.limiter = limiter
        self.granted = None
        self.password = tk.StringVar()
        self.option_add('*Font', ('Times New Roman', 12))
        panel = ttk.Frame(self, padding=22)
        panel.pack(fill='both', expand=True)
        ttk.Label(panel, text='5자리 비밀번호를 입력하세요.').pack(anchor='w', pady=(0, 14))
        ttk.Label(panel, text='사용기한: 2026-10-31까지').pack(anchor='w', pady=(0, 8))
        validate = self.register(lambda value: not value or (len(value) <= 5 and value.isascii() and value.isdigit()))
        self.entry = ttk.Entry(panel, textvariable=self.password, show='●', validate='key', validatecommand=(validate, '%P'))
        self.entry.pack(fill='x')
        ttk.Button(panel, text='실행', command=self.submit).pack(anchor='e', pady=16)
        self.bind('<Return>', lambda event: self.submit())
        self.bind('<Escape>', lambda event: self.destroy())
        self._focus_timer = self.after(100, self._focus_entry)

    def _focus_entry(self):
        self._focus_timer = None
        self.entry.focus_force()

    def destroy(self):
        if self._focus_timer is not None:
            self.after_cancel(self._focus_timer)
            self._focus_timer = None
        self.update_idletasks()
        super().destroy()

    def submit(self):
        try:
            used = self.limiter.authenticate(self.password.get())
        except LicenseExpired:
            messagebox.showinfo('사용기한 만료', EXPIRED_MESSAGE, parent=self)
            self.destroy()
            return
        except LaunchLimitReached:
            messagebox.showinfo('실행 제한', BLOCKED_MESSAGE, parent=self)
            self.destroy()
            return
        except AccessStateError as exc:
            messagebox.showerror('실행 기록 오류', str(exc), parent=self)
            self.destroy()
            return
        if used is None:
            messagebox.showerror('비밀번호 오류', '비밀번호가 올바르지 않습니다.', parent=self)
            self.password.set('')
            self.entry.focus_force()
            return
        self.granted = used
        self.destroy()


def main():
    limiter = LaunchLimiter()
    # Load and validate bundled history before a successful authentication can
    # consume a launch. A bad package must not use up the user's quota.
    try:
        ensure_not_expired()
        data = load_embedded_data()
        used = limiter.used()
        ensure_not_expired()
    except LicenseExpired:
        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo('사용기한 만료', EXPIRED_MESSAGE, parent=root)
        root.destroy()
        return
    except (ValueError, OSError, AccessStateError) as exc:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror('실행 오류', str(exc), parent=root)
        root.destroy()
        return
    if used >= 2:
        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo('실행 제한', BLOCKED_MESSAGE, parent=root)
        root.destroy()
        return
    login = PasswordWindow(limiter)
    login.mainloop()
    if login.granted is not None:
        try:
            app = OneTimeApp(data, login.granted)
        except LicenseExpired:
            root = tk.Tk()
            root.withdraw()
            messagebox.showinfo('사용기한 만료', EXPIRED_MESSAGE, parent=root)
            root.destroy()
            return
        app.mainloop()


if __name__ == '__main__':
    main()
