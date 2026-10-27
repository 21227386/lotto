"""Excel workbook input and output. Never evaluate workbook formulas."""
from pathlib import Path
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from openpyxl.utils.datetime import from_excel

from openpyxl import Workbook, load_workbook
from model import BASE_RATE, OUTPUT_SUBCOMPOSITION_PLAN, estimate, frequency_subgroups

HEADERS = ['Draw', 'Draw date', 'N1', 'N2', 'N3', 'N4', 'N5', 'N6']
LEGACY_HEADERS = ['회차', '번호1', '번호2', '번호3', '번호4', '번호5', '번호6']


@dataclass
class DrawInput:
    draws: list
    ids: list
    dates: list
    sheet: str
    source_count: int


def _integer(value, row_no):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not float(value).is_integer():
        raise ValueError(f'{row_no}행: 회차와 번호는 정수여야 합니다. 빈 셀·수식·문자를 제거하세요.')
    return int(value)


def _draw_date(value, row_no, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch)
            if isinstance(converted, datetime):
                return converted.date()
        except (ValueError, OverflowError):
            pass
    if isinstance(value, str):
        for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d'):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError(f'{row_no}행: Draw date에 엑셀 날짜 또는 YYYY-MM-DD 형식의 추첨일을 넣으세요.')


def read_workbook(path):
    if Path(path).suffix.lower() != '.xlsx':
        raise ValueError('Excel 통합 문서(.xlsx)를 선택하세요. .xls는 Excel에서 .xlsx로 저장하세요.')
    try:
        book = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError('엑셀 파일을 읽을 수 없습니다. 암호가 없는 정상 .xlsx 파일인지 확인하세요.') from exc
    try:
        sheet = book['Lotto Draws'] if 'Lotto Draws' in book.sheetnames else book.active
        records = []
        layout = None
        seen_ids = set()
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            while values and values[-1] is None:
                values.pop()
            if not values:
                continue
            row_no = row[0].row
            if layout is None:
                header = [str(v).strip().lower() for v in values]
                if header == [v.lower() for v in HEADERS]:
                    layout = 8
                    continue
                legacy = [LEGACY_HEADERS, LEGACY_HEADERS[1:],
                          ['draw'] + [f'n{i}' for i in range(1, 7)], [f'n{i}' for i in range(1, 7)]]
                for item in legacy:
                    if header == [v.lower() for v in item]:
                        layout = len(item)
                        break
                if layout is not None:
                    continue
                if len(values) not in (6, 7):
                    raise ValueError('8열 입력은 첫 행에 Draw, Draw date, N1~N6 헤더가 필요합니다.')
                layout = len(values)
            if len(values) != layout:
                raise ValueError(f'{row_no}행: 열 개수가 다릅니다. 빈 셀·보너스 번호·추가 열을 확인하세요.')
            draw_id = _integer(values[0], row_no) if layout >= 7 else None
            draw_date = _draw_date(values[1], row_no, book.epoch) if layout == 8 else None
            numbers = tuple(_integer(v, row_no) for v in values[2 if layout == 8 else 1 if layout == 7 else 0:])
            if len(set(numbers)) != 6 or any(n < 1 or n > 45 for n in numbers):
                raise ValueError(f'{row_no}행: 1~45의 서로 다른 본번호 6개가 필요합니다.')
            if draw_id is not None:
                if draw_id <= 0 or draw_id in seen_ids:
                    raise ValueError(f'{row_no}행: 회차는 중복 없는 양의 정수여야 합니다.')
                seen_ids.add(draw_id)
            records.append((draw_id, draw_date, tuple(sorted(numbers))))
        source_count = len(records)
        if source_count < 412:
            raise ValueError(f'최소 412회가 필요합니다. 현재 {source_count}회입니다.')
        if layout >= 7:
            records.sort(key=lambda item: item[0])
        selected = records
        ids = [item[0] for item in selected] if layout >= 7 else []
        dates = [item[1] for item in selected] if layout == 8 else []
        if ids and ids != list(range(ids[0], ids[-1] + 1)):
            raise ValueError('입력 데이터에 누락 회차가 있습니다. Draw 열을 확인하세요.')
        if dates and any(b <= a for a, b in zip(dates, dates[1:])):
            raise ValueError('회차 순서와 추첨일 순서가 맞지 않습니다. Draw date를 확인하세요.')
        return DrawInput([item[2] for item in selected], ids, dates, sheet.title, source_count)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('엑셀 시트 데이터를 읽을 수 없습니다. 파일을 다시 저장한 후 열어 주세요.') from exc
    finally:
        book.close()


def read_draws(path):
    data = read_workbook(path)
    return data.draws, data.ids, data.sheet


def write_template(path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Lotto Draws'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'C2'
    for column in 'ABCDEFGH':
        sheet.column_dimensions[column].width = 16
    for row in range(2, 414):
        sheet.cell(row, 2).number_format = 'yyyy-mm-dd'
    sheet.auto_filter.ref = 'A1:H413'
    book.save(path)
    book.close()


def write_analysis(path, estimates, numbers, ids, strength, uncertainty, dates=None, source_count=None):
    book = Workbook()
    summary = book.active
    summary.title = '추천결과'
    summary.append(['항목', '값'])
    summary.append(['대상 회차', max(ids) + 1 if ids else '회차 미확인'])
    summary.append(['입력 범위', f'{min(ids)}~{max(ids)}' if ids else '회차 없음'])
    summary.append(['사전 강도', strength])
    summary.append(['사후 불확실성 반영', '예' if uncertainty else '아니오'])
    summary.append(['추천 번호'] + list(numbers))
    summary.append(['모델', 'Beta–Binomial 주변 사후근사 + 순차 가중 비복원추출'])
    summary.append(['주의', '추정 포함확률은 최종 추출 포함확률과 다름. 당첨 우위 미입증.'])
    if dates:
        summary.append(['최신 추첨일', dates[-1].isoformat()])
        summary.append(['다음 추첨 예정일', (dates[-1] + timedelta(days=7)).isoformat()])
    if source_count is not None:
        summary.append(['파일 전체 회차 수', source_count])
        summary.append(['분석 회차 수', sum(e.count for e in estimates) // 6])
    sheet = book.create_sheet('번호별분석')
    sheet.append(['번호', '출현횟수', '사후 포함확률 추정', '균등 대비(%p)', '사후 표준편차(%p)'])
    for e in estimates:
        sheet.append([e.number, e.count, e.mean, 100 * (e.mean - BASE_RATE), 100 * e.sd])
        sheet.cell(sheet.max_row, 3).number_format = '0.000%'
    sheet.freeze_panes = 'A2'
    for ws in book:
        for col in 'ABCDEFG':
            ws.column_dimensions[col].width = 26
    book.save(path)
    book.close()


def write_recommendations(path, recommendations, data, strength):
    book = Workbook()
    sheet = book.active
    sheet.title = '추천번호'
    sheet.append(['Set', 'N1', 'N2', 'N3', 'N4', 'N5', 'N6', '고:중:저', '세부 구성(고, 중, 저 순)'])
    for index, result in enumerate(recommendations, 1):
        sheet.append([index] + list(result.numbers) + [':'.join(map(str, result.composition)),
                     ', '.join(':'.join(map(str, part)) for part in result.subcomposition)])
    sheet.freeze_panes = 'B2'
    meta = book.create_sheet('설정')
    meta.append(['항목', '값'])
    meta.append(['대상 회차', data.ids[-1] + 1 if data.ids else '회차 미확인'])
    meta.append(['분석 회차 수', len(data.draws)])
    meta.append(['분석 범위', f'{data.ids[0]}~{data.ids[-1]}' if data.ids else '회차 없음'])
    meta.append(['사전 강도', strength])
    meta.append(['순위 기준', '9개 세부 빈도 그룹 구성별 베이지안 가중치 PL 조합 확률'])
    meta.append(['선택 방식', '같은 세부 그룹 안에서 동일 가중치인 번호만 무작위 선택, 세트 간 반복 허용'])
    meta.append(['결과 유지', '반복 클릭은 기존 결과 유지, 파일 재선택 시 새로 추첨'])
    meta.append(['주의', '모델 기준 상위 조합이며 실제 당첨 우위는 입증되지 않았습니다.'])
    meta.append(['고:중:저 구성', '기존 5가지 세부 구성마다 상위 2세트, 총 10세트'])
    meta.append(['번호 제한', '연속 3개 금지 / n,n+7,n+14 금지'])
    meta.append(['공개 순서', '세트 안에서 베이지안 가중치 높은 번호부터'])
    meta.append(['분류 기준', '출현 횟수 내림차순 각 15개, 동률은 번호 오름차순'])
    meta.append(['세부 분류', '각 15개를 베이지안 가중치 순 상·중·하 5개씩 분류'])
    for index, plan in enumerate(OUTPUT_SUBCOMPOSITION_PLAN, 1):
        meta.append([f'SET {index} 세부 구성', ', '.join(':'.join(map(str, part)) for part in plan)])
    estimates = estimate(data.draws, strength)
    groups = frequency_subgroups(estimates)
    group_sheet = book.create_sheet('빈도그룹')
    group_sheet.append(['빈도 그룹', '세부 그룹', '번호', '출현 횟수', '베이지안 가중치'])
    lookup = {e.number: e for e in estimates}
    for name, group in zip(['고빈도', '중빈도', '저빈도'], groups):
        for subname, part in zip(['상위', '중위', '하위'], group):
            for number in part:
                group_sheet.append([name, subname, number, lookup[number].count, lookup[number].mean])
    if data.dates:
        meta.append(['최신 추첨일', data.dates[-1].isoformat()])
        meta.append(['다음 추첨 예정일', (data.dates[-1] + timedelta(days=7)).isoformat()])
    for ws in book:
        for col in 'ABCDEFGHI':
            ws.column_dimensions[col].width = 18 if ws == sheet else 28
    book.save(path)
    book.close()
